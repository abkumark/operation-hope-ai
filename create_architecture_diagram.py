#!/usr/bin/env python3
"""
Generate Architecture Diagram for Operation HOPE AI System

Creates a professional PowerPoint slide showing the system architecture,
components, data flow, and technology stack.
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR_TYPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_architecture_diagram():
    """Create a comprehensive architecture diagram slide."""

    # Initialize presentation with widescreen layout
    prs = Presentation()
    prs.slide_width = Inches(13.33)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Use blank slide layout
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Brand colors (Operation HOPE)
    BRAND_PRIMARY = RGBColor(27, 160, 215)    # Blue
    BRAND_SECONDARY = RGBColor(0, 80, 115)    # Dark blue
    BRAND_ACCENT = RGBColor(255, 165, 0)      # Orange
    BRAND_SUCCESS = RGBColor(76, 175, 80)     # Green
    BRAND_NEUTRAL = RGBColor(242, 242, 242)   # Light gray
    BRAND_TEXT = RGBColor(51, 51, 51)         # Dark gray

    # Layout dimensions
    slide_width = 13.33
    slide_height = 7.5
    margin = 0.5

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(margin), Inches(0.2),
        Inches(slide_width - 2 * margin), Inches(0.8)
    )
    title_frame = title_box.text_frame
    title_frame.text = "Operation HOPE AI - System Architecture"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = BRAND_SECONDARY
    title_para.alignment = PP_ALIGN.CENTER

    # Subtitle
    subtitle_box = slide.shapes.add_textbox(
        Inches(margin), Inches(0.9),
        Inches(slide_width - 2 * margin), Inches(0.4)
    )
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "AI-Powered Support Ticket System • ~30 tickets/day • 48 hours → minutes response time"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(14)
    subtitle_para.font.color.rgb = BRAND_TEXT
    subtitle_para.alignment = PP_ALIGN.CENTER

    # Layer positions
    y_layers = {
        'users': 1.6,
        'frontend': 2.7,
        'api': 3.8,
        'ai': 4.9,
        'storage': 6.0
    }

    # Component dimensions
    box_width = 2.2
    box_height = 0.7
    small_box_width = 1.8
    small_box_height = 0.5

    # === USER LAYER ===
    users_y = y_layers['users']

    # Public Users
    public_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1), Inches(users_y),
        Inches(small_box_width), Inches(small_box_height)
    )
    public_shape.fill.solid()
    public_shape.fill.fore_color.rgb = BRAND_ACCENT
    public_shape.line.color.rgb = BRAND_SECONDARY
    public_shape.line.width = Pt(2)
    public_text = public_shape.text_frame
    public_text.text = "Public Users\n(Ticket Submitters)"
    public_text.paragraphs[0].font.size = Pt(10)
    public_text.paragraphs[0].font.bold = True
    public_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    public_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Engineers
    engineers_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4.5), Inches(users_y),
        Inches(small_box_width), Inches(small_box_height)
    )
    engineers_shape.fill.solid()
    engineers_shape.fill.fore_color.rgb = BRAND_PRIMARY
    engineers_shape.line.color.rgb = BRAND_SECONDARY
    engineers_shape.line.width = Pt(2)
    engineers_text = engineers_shape.text_frame
    engineers_text.text = "Engineers\n(5 Support Staff)"
    engineers_text.paragraphs[0].font.size = Pt(10)
    engineers_text.paragraphs[0].font.bold = True
    engineers_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    engineers_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Admin
    admin_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8), Inches(users_y),
        Inches(small_box_width), Inches(small_box_height)
    )
    admin_shape.fill.solid()
    admin_shape.fill.fore_color.rgb = BRAND_SECONDARY
    admin_shape.line.color.rgb = BRAND_SECONDARY
    admin_shape.line.width = Pt(2)
    admin_text = admin_shape.text_frame
    admin_text.text = "Admin\n(Analytics & Config)"
    admin_text.paragraphs[0].font.size = Pt(10)
    admin_text.paragraphs[0].font.bold = True
    admin_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    admin_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # === FRONTEND LAYER ===
    frontend_y = y_layers['frontend']

    frontend_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4), Inches(frontend_y),
        Inches(box_width * 2), Inches(box_height)
    )
    frontend_shape.fill.solid()
    frontend_shape.fill.fore_color.rgb = BRAND_NEUTRAL
    frontend_shape.line.color.rgb = BRAND_TEXT
    frontend_shape.line.width = Pt(2)
    frontend_text = frontend_shape.text_frame
    frontend_text.text = "Frontend - Single Page App\nHTML/CSS/JS • Chart.js • Vanilla JavaScript"
    frontend_text.paragraphs[0].font.size = Pt(12)
    frontend_text.paragraphs[0].font.bold = True
    frontend_text.paragraphs[0].font.color.rgb = BRAND_TEXT
    frontend_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # === API LAYER ===
    api_y = y_layers['api']

    api_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4), Inches(api_y),
        Inches(box_width * 2), Inches(box_height)
    )
    api_shape.fill.solid()
    api_shape.fill.fore_color.rgb = BRAND_PRIMARY
    api_shape.line.color.rgb = BRAND_SECONDARY
    api_shape.line.width = Pt(2)
    api_text = api_shape.text_frame
    api_text.text = "FastAPI Backend\n37 Endpoints • Authentication/RBAC • Route Management"
    api_text.paragraphs[0].font.size = Pt(12)
    api_text.paragraphs[0].font.bold = True
    api_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    api_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # === AI PROCESSING LAYER ===
    ai_y = y_layers['ai']

    # Classification
    classify_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1), Inches(ai_y),
        Inches(box_width), Inches(box_height)
    )
    classify_shape.fill.solid()
    classify_shape.fill.fore_color.rgb = BRAND_ACCENT
    classify_shape.line.color.rgb = BRAND_SECONDARY
    classify_shape.line.width = Pt(2)
    classify_text = classify_shape.text_frame
    classify_text.text = "AI Classification\nGPT-5.2 Azure OpenAI\n21 Categories"
    classify_text.paragraphs[0].font.size = Pt(10)
    classify_text.paragraphs[0].font.bold = True
    classify_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    classify_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Routing
    routing_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(4), Inches(ai_y),
        Inches(box_width), Inches(box_height)
    )
    routing_shape.fill.solid()
    routing_shape.fill.fore_color.rgb = BRAND_ACCENT
    routing_shape.line.color.rgb = BRAND_SECONDARY
    routing_shape.line.width = Pt(2)
    routing_text = routing_shape.text_frame
    routing_text.text = "Smart Routing\nConfidence-Based\nEngineer Matching"
    routing_text.paragraphs[0].font.size = Pt(10)
    routing_text.paragraphs[0].font.bold = True
    routing_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    routing_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # RAG Response
    rag_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7), Inches(ai_y),
        Inches(box_width), Inches(box_height)
    )
    rag_shape.fill.solid()
    rag_shape.fill.fore_color.rgb = BRAND_ACCENT
    rag_shape.line.color.rgb = BRAND_SECONDARY
    rag_shape.line.width = Pt(2)
    rag_text = rag_shape.text_frame
    rag_text.text = "RAG Response\nKnowledge Base\nAuto-Resolution"
    rag_text.paragraphs[0].font.size = Pt(10)
    rag_text.paragraphs[0].font.bold = True
    rag_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    rag_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # === STORAGE LAYER ===
    storage_y = y_layers['storage']

    # SQLite Database
    sqlite_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(2), Inches(storage_y),
        Inches(box_width), Inches(box_height)
    )
    sqlite_shape.fill.solid()
    sqlite_shape.fill.fore_color.rgb = BRAND_SUCCESS
    sqlite_shape.line.color.rgb = BRAND_SECONDARY
    sqlite_shape.line.width = Pt(2)
    sqlite_text = sqlite_shape.text_frame
    sqlite_text.text = "SQLite Database\nTickets • Analytics\nUser Management"
    sqlite_text.paragraphs[0].font.size = Pt(10)
    sqlite_text.paragraphs[0].font.bold = True
    sqlite_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    sqlite_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # ChromaDB Vector Store
    chroma_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5), Inches(storage_y),
        Inches(box_width), Inches(box_height)
    )
    chroma_shape.fill.solid()
    chroma_shape.fill.fore_color.rgb = BRAND_SUCCESS
    chroma_shape.line.color.rgb = BRAND_SECONDARY
    chroma_shape.line.width = Pt(2)
    chroma_text = chroma_shape.text_frame
    chroma_text.text = "ChromaDB\nVector Embeddings\n24+ KB Articles"
    chroma_text.paragraphs[0].font.size = Pt(10)
    chroma_text.paragraphs[0].font.bold = True
    chroma_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    chroma_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # === EXTERNAL INTEGRATIONS ===

    # SMTP Notifications
    smtp_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.5), Inches(3.5),
        Inches(1.5), Inches(0.5)
    )
    smtp_shape.fill.solid()
    smtp_shape.fill.fore_color.rgb = BRAND_SECONDARY
    smtp_shape.line.color.rgb = BRAND_SECONDARY
    smtp_shape.line.width = Pt(1)
    smtp_text = smtp_shape.text_frame
    smtp_text.text = "SMTP\nNotifications"
    smtp_text.paragraphs[0].font.size = Pt(9)
    smtp_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    smtp_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Dynamics 365
    dynamics_shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(10.5), Inches(4.5),
        Inches(1.5), Inches(0.5)
    )
    dynamics_shape.fill.solid()
    dynamics_shape.fill.fore_color.rgb = BRAND_SECONDARY
    dynamics_shape.line.color.rgb = BRAND_SECONDARY
    dynamics_shape.line.width = Pt(1)
    dynamics_text = dynamics_shape.text_frame
    dynamics_text.text = "Dynamics 365\n(Optional)"
    dynamics_text.paragraphs[0].font.size = Pt(9)
    dynamics_text.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    dynamics_text.paragraphs[0].alignment = PP_ALIGN.CENTER

    # === DATA FLOW ARROWS ===

    # User to Frontend
    arrow1 = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(5.4), Inches(users_y + small_box_height),
        Inches(5.4), Inches(frontend_y)
    )
    arrow1.line.color.rgb = BRAND_TEXT
    arrow1.line.width = Pt(3)

    # Frontend to API
    arrow2 = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(5.4), Inches(frontend_y + box_height),
        Inches(5.4), Inches(api_y)
    )
    arrow2.line.color.rgb = BRAND_TEXT
    arrow2.line.width = Pt(3)

    # API to AI Pipeline (Classification)
    arrow3 = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(4.5), Inches(api_y + box_height),
        Inches(2.1), Inches(ai_y)
    )
    arrow3.line.color.rgb = BRAND_TEXT
    arrow3.line.width = Pt(3)

    # Classification to Routing
    arrow4 = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(3.2), Inches(ai_y + box_height/2),
        Inches(4), Inches(ai_y + box_height/2)
    )
    arrow4.line.color.rgb = BRAND_ACCENT
    arrow4.line.width = Pt(3)

    # Routing to RAG
    arrow5 = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(6.2), Inches(ai_y + box_height/2),
        Inches(7), Inches(ai_y + box_height/2)
    )
    arrow5.line.color.rgb = BRAND_ACCENT
    arrow5.line.width = Pt(3)

    # AI to Storage
    arrow6 = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        Inches(5.4), Inches(ai_y + box_height),
        Inches(5.4), Inches(storage_y)
    )
    arrow6.line.color.rgb = BRAND_TEXT
    arrow6.line.width = Pt(3)

    # === LEGEND ===
    legend_x = 10.5
    legend_y = 1.5

    legend_title = slide.shapes.add_textbox(
        Inches(legend_x), Inches(legend_y),
        Inches(2), Inches(0.3)
    )
    legend_title_frame = legend_title.text_frame
    legend_title_frame.text = "Legend"
    legend_title_para = legend_title_frame.paragraphs[0]
    legend_title_para.font.size = Pt(12)
    legend_title_para.font.bold = True
    legend_title_para.font.color.rgb = BRAND_TEXT

    # Legend items
    legend_items = [
        ("AI Processing", BRAND_ACCENT),
        ("Storage Layer", BRAND_SUCCESS),
        ("User Interface", BRAND_NEUTRAL),
        ("API Layer", BRAND_PRIMARY)
    ]

    for i, (label, color) in enumerate(legend_items):
        y_pos = legend_y + 0.4 + (i * 0.25)

        # Color box
        color_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(legend_x), Inches(y_pos),
            Inches(0.15), Inches(0.15)
        )
        color_box.fill.solid()
        color_box.fill.fore_color.rgb = color
        color_box.line.color.rgb = BRAND_TEXT

        # Label
        label_box = slide.shapes.add_textbox(
            Inches(legend_x + 0.2), Inches(y_pos),
            Inches(1.5), Inches(0.2)
        )
        label_frame = label_box.text_frame
        label_frame.text = label
        label_para = label_frame.paragraphs[0]
        label_para.font.size = Pt(9)
        label_para.font.color.rgb = BRAND_TEXT

    # Save presentation
    output_path = "Operation_HOPE_AI_Architecture.pptx"
    prs.save(output_path)
    print(f"Architecture diagram saved as: {output_path}")

    return output_path


if __name__ == "__main__":
    create_architecture_diagram()